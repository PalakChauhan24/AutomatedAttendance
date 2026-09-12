import tkinter as tk
from tkinter import ttk, messagebox
import csv
import os
import subprocess
import sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ==================================================
# PATHS
# ==================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

STUDENTS_FILE = os.path.join(
    BASE_DIR,
    "data",
    "students.csv"
)

ATTENDANCE_FOLDER = os.path.join(
    BASE_DIR,
    "attendance"
)

ATTENDANCE_RECORDS_FILE = os.path.join(
    ATTENDANCE_FOLDER,
    "attendance_records.xlsx"
)

REPORT_FILE = os.path.join(
    ATTENDANCE_FOLDER,
    "attendance_report.xlsx"
)

os.makedirs(
    ATTENDANCE_FOLDER,
    exist_ok=True
)


# ==================================================
# COLORS
# ==================================================

BG = "#0F172A"
CARD = "#1E293B"
WHITE = "#F8FAFC"
LIGHT_TEXT = "#CBD5E1"
MUTED = "#94A3B8"
BLUE = "#3B82F6"
BLUE_HOVER = "#2563EB"
GREEN = "#22C55E"
ORANGE = "#F59E0B"


# ==================================================
# HEADERS
# ==================================================

STUDENT_HEADERS = [
    "ID",
    "Name",
    "Enrollment No",
    "Email",
    "Year",
    "Semester",
    "Section"
]

ATTENDANCE_HEADERS = [
    "ID",
    "Name",
    "Enrollment No",
    "Year",
    "Semester",
    "Section",
    "Date",
    "Time",
    "Status"
]


# ==================================================
# LOAD STUDENTS
# ==================================================

students = {}


def load_students():

    global students

    students = {}

    if not os.path.exists(
        STUDENTS_FILE
    ):
        return

    try:

        with open(
            STUDENTS_FILE,
            "r",
            newline="",
            encoding="utf-8-sig"
        ) as file:

            reader = csv.DictReader(file)

            for row in reader:

                try:

                    student_id = int(
                        str(
                            row.get("ID", "")
                        ).strip()
                    )

                except Exception:
                    continue

                students[student_id] = {
                    "name": str(
                        row.get("Name", "")
                    ).strip(),

                    "enrollment": str(
                        row.get("Enrollment No", "")
                    ).strip(),

                    "email": str(
                        row.get("Email", "")
                    ).strip(),

                    "year": str(
                        row.get(
                            "Year",
                            "Not Specified"
                        )
                    ).strip(),

                    "semester": str(
                        row.get(
                            "Semester",
                            "Not Specified"
                        )
                    ).strip(),

                    "section": str(
                        row.get(
                            "Section",
                            "Not Specified"
                        )
                    ).strip()
                }

    except Exception:
        pass


load_students()


# ==================================================
# INITIALIZE / UPGRADE EXCEL
# ==================================================

def initialize_attendance_excel():

    if not os.path.exists(
        ATTENDANCE_RECORDS_FILE
    ):

        workbook = Workbook()

        attendance_sheet = (
            workbook.active
        )

        attendance_sheet.title = (
            "Attendance Records"
        )

        attendance_sheet.append(
            ATTENDANCE_HEADERS
        )

        student_sheet = (
            workbook.create_sheet(
                "Student Records"
            )
        )

        student_sheet.append(
            STUDENT_HEADERS
        )

        workbook.save(
            ATTENDANCE_RECORDS_FILE
        )

    else:

        upgrade_excel()

    update_student_records()


def upgrade_excel():

    try:

        workbook = load_workbook(
            ATTENDANCE_RECORDS_FILE
        )

        # ==================================================
        # ATTENDANCE RECORDS
        # ==================================================

        if "Attendance Records" not in (
            workbook.sheetnames
        ):

            sheet = workbook.create_sheet(
                "Attendance Records",
                0
            )

            sheet.append(
                ATTENDANCE_HEADERS
            )

        else:

            sheet = workbook[
                "Attendance Records"
            ]

            current_headers = [
                sheet.cell(
                    1,
                    column
                ).value
                for column in range(
                    1,
                    sheet.max_column + 1
                )
            ]

            if current_headers != ATTENDANCE_HEADERS:

                old_rows = list(
                    sheet.iter_rows(
                        min_row=2,
                        values_only=True
                    )
                )

                workbook.remove(
                    sheet
                )

                sheet = workbook.create_sheet(
                    "Attendance Records",
                    0
                )

                sheet.append(
                    ATTENDANCE_HEADERS
                )

                index = {
                    str(header): position
                    for position, header
                    in enumerate(
                        current_headers
                    )
                }

                for old_row in old_rows:

                    def value(header):

                        position = index.get(
                            header
                        )

                        if position is None:
                            return ""

                        if position >= len(
                            old_row
                        ):
                            return ""

                        return (
                            old_row[position]
                            or ""
                        )

                    sheet.append([
                        value("ID"),
                        value("Name"),
                        value("Enrollment No"),
                        value("Year")
                        or "Not Specified",
                        value("Semester")
                        or "Not Specified",
                        value("Section")
                        or "Not Specified",
                        value("Date"),
                        value("Time"),
                        value("Status")
                    ])

        # ==================================================
        # STUDENT RECORDS
        # ==================================================

        if "Student Records" not in (
            workbook.sheetnames
        ):

            student_sheet = (
                workbook.create_sheet(
                    "Student Records"
                )
            )

            student_sheet.append(
                STUDENT_HEADERS
            )

        else:

            student_sheet = (
                workbook["Student Records"]
            )

            current_headers = [
                student_sheet.cell(
                    1,
                    column
                ).value
                for column in range(
                    1,
                    student_sheet.max_column + 1
                )
            ]

            if current_headers != STUDENT_HEADERS:

                old_rows = list(
                    student_sheet.iter_rows(
                        min_row=2,
                        values_only=True
                    )
                )

                workbook.remove(
                    student_sheet
                )

                student_sheet = (
                    workbook.create_sheet(
                        "Student Records"
                    )
                )

                student_sheet.append(
                    STUDENT_HEADERS
                )

                index = {
                    str(header): position
                    for position, header
                    in enumerate(
                        current_headers
                    )
                }

                for old_row in old_rows:

                    def value(header):

                        position = index.get(
                            header
                        )

                        if position is None:
                            return ""

                        if position >= len(
                            old_row
                        ):
                            return ""

                        return (
                            old_row[position]
                            or ""
                        )

                    student_sheet.append([
                        value("ID"),
                        value("Name"),
                        value("Enrollment No"),
                        value("Email"),
                        value("Year")
                        or "Not Specified",
                        value("Semester")
                        or "Not Specified",
                        value("Section")
                        or "Not Specified"
                    ])

        workbook.save(
            ATTENDANCE_RECORDS_FILE
        )

    except Exception:
        pass


# ==================================================
# UPDATE STUDENT RECORDS IN EXCEL
# ==================================================

def update_student_records():

    if not os.path.exists(
        ATTENDANCE_RECORDS_FILE
    ):
        return

    try:

        workbook = load_workbook(
            ATTENDANCE_RECORDS_FILE
        )

        if "Student Records" not in (
            workbook.sheetnames
        ):

            sheet = workbook.create_sheet(
                "Student Records"
            )

            sheet.append(
                STUDENT_HEADERS
            )

        else:

            sheet = workbook[
                "Student Records"
            ]

        if sheet.max_row > 1:

            sheet.delete_rows(
                2,
                sheet.max_row - 1
            )

        for student_id in sorted(
            students.keys()
        ):

            student = students[
                student_id
            ]

            sheet.append([
                student_id,
                student["name"],
                student["enrollment"],
                student["email"],
                student["year"],
                student["semester"],
                student["section"]
            ])

        widths = {
            "A": 10,
            "B": 25,
            "C": 25,
            "D": 35,
            "E": 12,
            "F": 14,
            "G": 12
        }

        for column, width in widths.items():

            sheet.column_dimensions[
                column
            ].width = width

        workbook.save(
            ATTENDANCE_RECORDS_FILE
        )

    except Exception:
        pass


# ==================================================
# READ EXCEL ATTENDANCE
# ==================================================

def get_excel_attendance_records():

    records = []

    if not os.path.exists(
        ATTENDANCE_RECORDS_FILE
    ):
        return records

    try:

        workbook = load_workbook(
            ATTENDANCE_RECORDS_FILE,
            read_only=True
        )

        if "Attendance Records" not in (
            workbook.sheetnames
        ):

            workbook.close()
            return records

        sheet = workbook[
            "Attendance Records"
        ]

        headers = [
            cell.value
            for cell in sheet[1]
        ]

        header_index = {
            str(header): index
            for index, header
            in enumerate(headers)
            if header is not None
        }

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True
        ):

            if not row:
                continue

            def value(header):
                position = header_index.get(
                    header
                )

                if position is None:
                    return ""

                if position >= len(row):
                    return ""

                return row[position]

            if value("ID") is None:
                continue

            records.append({
                "ID": str(
                    value("ID") or ""
                ).strip(),

                "Name": str(
                    value("Name") or ""
                ).strip(),

                "Enrollment No": str(
                    value("Enrollment No") or ""
                ).strip(),

                "Year": str(
                    value("Year")
                    or "Not Specified"
                ).strip(),

                "Semester": str(
                    value("Semester")
                    or "Not Specified"
                ).strip(),

                "Section": str(
                    value("Section")
                    or "Not Specified"
                ).strip(),

                "Date": str(
                    value("Date") or ""
                ).strip(),

                "Time": str(
                    value("Time") or ""
                ).strip(),

                "Status": str(
                    value("Status") or ""
                ).strip()
            })

        workbook.close()

    except Exception:
        return []

    return records


# ==================================================
# CSV FALLBACK
# ==================================================

def get_csv_attendance_records():

    records = []

    if not os.path.exists(
        ATTENDANCE_FOLDER
    ):
        return records

    try:
        filenames = os.listdir(
            ATTENDANCE_FOLDER
        )
    except Exception:
        return records

    for filename in sorted(
        filenames
    ):

        if not (
            filename.startswith(
                "attendance_"
            )
            and filename.endswith(".csv")
        ):
            continue

        file_path = os.path.join(
            ATTENDANCE_FOLDER,
            filename
        )

        try:

            with open(
                file_path,
                "r",
                newline="",
                encoding="utf-8-sig"
            ) as file:

                reader = csv.DictReader(
                    file
                )

                for row in reader:

                    records.append({
                        "ID": str(
                            row.get("ID", "")
                        ).strip(),

                        "Name": str(
                            row.get("Name", "")
                        ).strip(),

                        "Enrollment No": str(
                            row.get(
                                "Enrollment No",
                                ""
                            )
                        ).strip(),

                        "Year": str(
                            row.get(
                                "Year",
                                "Not Specified"
                            )
                        ).strip(),

                        "Semester": str(
                            row.get(
                                "Semester",
                                "Not Specified"
                            )
                        ).strip(),

                        "Section": str(
                            row.get(
                                "Section",
                                "Not Specified"
                            )
                        ).strip(),

                        "Date": str(
                            row.get("Date", "")
                        ).strip(),

                        "Time": str(
                            row.get("Time", "")
                        ).strip(),

                        "Status": str(
                            row.get("Status", "")
                        ).strip()
                    })

        except Exception:
            continue

    return records


# ==================================================
# MIGRATE OLD CSV RECORDS INTO EXCEL
# ==================================================

def migrate_csv_to_excel():

    excel_records = (
        get_excel_attendance_records()
    )

    existing_keys = set()

    for record in excel_records:

        existing_keys.add((
            record["ID"],
            record["Date"]
        ))

    csv_records = (
        get_csv_attendance_records()
    )

    if not csv_records:
        return

    try:

        workbook = load_workbook(
            ATTENDANCE_RECORDS_FILE
        )

        sheet = workbook[
            "Attendance Records"
        ]

        changed = False

        for record in csv_records:

            key = (
                record["ID"],
                record["Date"]
            )

            if key in existing_keys:
                continue

            sheet.append([
                record["ID"],
                record["Name"],
                record["Enrollment No"],
                record["Year"],
                record["Semester"],
                record["Section"],
                record["Date"],
                record["Time"],
                record["Status"]
            ])

            existing_keys.add(key)
            changed = True

        if changed:

            workbook.save(
                ATTENDANCE_RECORDS_FILE
            )

    except Exception:
        pass


# ==================================================
# GET ALL ATTENDANCE RECORDS
# ==================================================

def get_attendance_records():

    initialize_attendance_excel()

    migrate_csv_to_excel()

    records = (
        get_excel_attendance_records()
    )

    if records:
        return records

    return get_csv_attendance_records()


# ==================================================
# ATTENDANCE SUMMARY DATA
# ==================================================

def get_attendance_data():

    records = get_attendance_records()

    present_count = {
        student_id: 0
        for student_id in students
    }

    attendance_dates = set()
    student_day_pairs = set()

    for record in records:

        try:

            student_id = int(
                record["ID"]
            )

        except Exception:
            continue

        date = record["Date"]

        status = (
            record["Status"]
            .strip()
            .lower()
        )

        if date:
            attendance_dates.add(
                date
            )

        if status == "present":

            pair = (
                student_id,
                date
            )

            if pair not in student_day_pairs:

                student_day_pairs.add(
                    pair
                )

                if student_id in present_count:

                    present_count[
                        student_id
                    ] += 1

    return (
        records,
        len(attendance_dates),
        present_count
    )


# ==================================================
# CLEAR TABLE
# ==================================================

def clear_table():

    for item in table.get_children():
        table.delete(item)


# ==================================================
# SUMMARY HEADINGS
# ==================================================

def set_summary_headings():

    headings = (
        "ID",
        "Name",
        "Enrollment No",
        "Year",
        "Semester",
        "Section",
        "Present",
        "Total Days",
        "Attendance %"
    )

    for column, heading in zip(
        columns,
        headings
    ):

        table.heading(
            column,
            text=heading
        )


# ==================================================
# DAILY HEADINGS
# ==================================================

def set_daily_headings():

    headings = (
        "ID",
        "Name",
        "Enrollment No",
        "Year",
        "Semester",
        "Section",
        "Date",
        "Time",
        "Status"
    )

    for column, heading in zip(
        columns,
        headings
    ):

        table.heading(
            column,
            text=heading
        )


# ==================================================
# SHOW SUMMARY
# ==================================================

def calculate_attendance():

    set_summary_headings()

    (
        records,
        total_days,
        present_count
    ) = get_attendance_data()

    clear_table()

    if total_days == 0:

        summary_label.config(
            text="No attendance records available"
        )

        status_label.config(
            text="No attendance data found",
            fg=ORANGE
        )

        return

    for student_id, student in (
        students.items()
    ):

        present = present_count.get(
            student_id,
            0
        )

        percentage = (
            present / total_days
        ) * 100

        table.insert(
            "",
            tk.END,
            values=(
                student_id,
                student["name"],
                student["enrollment"],
                student["year"],
                student["semester"],
                student["section"],
                present,
                total_days,
                f"{percentage:.2f}%"
            )
        )

    summary_label.config(
        text=(
            f"  Students: {len(students)}     "
            f"Attendance Days: {total_days}     "
            "Storage: Excel"
        )
    )

    status_label.config(
        text="Showing attendance summary",
        fg=GREEN
    )


# ==================================================
# SHOW DAILY RECORDS
# ==================================================

def show_daily_attendance():

    set_daily_headings()

    records = get_attendance_records()

    clear_table()

    if not records:

        summary_label.config(
            text="No attendance records available"
        )

        status_label.config(
            text="No attendance data found",
            fg=ORANGE
        )

        return

    for record in records:

        table.insert(
            "",
            tk.END,
            values=(
                record["ID"],
                record["Name"],
                record["Enrollment No"],
                record["Year"],
                record["Semester"],
                record["Section"],
                record["Date"],
                record["Time"],
                record["Status"]
            )
        )

    summary_label.config(
        text=(
            f"  Attendance Records: "
            f"{len(records)}     "
            "Storage: Excel"
        )
    )

    status_label.config(
        text="Showing daily attendance records",
        fg=GREEN
    )


# ==================================================
# OPEN ANALYTICS
# ==================================================

def open_analytics():

    analytics_file = os.path.join(
        BASE_DIR,
        "analytics.py"
    )

    if not os.path.exists(
        analytics_file
    ):

        messagebox.showerror(
            "Analytics Not Found",
            "analytics.py was not found."
        )

        return

    try:

        subprocess.Popen([
            sys.executable,
            analytics_file
        ])

        status_label.config(
            text="Analytics opened",
            fg=GREEN
        )

    except Exception as error:

        messagebox.showerror(
            "Analytics Error",
            f"Could not open analytics.\n\n{error}"
        )


# ==================================================
# EXPORT REPORT
# ==================================================

def export_to_excel():

    (
        records,
        total_days,
        present_count
    ) = get_attendance_data()

    if not records:

        messagebox.showwarning(
            "No Attendance",
            "There are no attendance records to export."
        )

        return

    try:

        workbook = Workbook()

        # ==================================================
        # SUMMARY
        # ==================================================

        summary_sheet = (
            workbook.active
        )

        summary_sheet.title = (
            "Attendance Summary"
        )

        summary_headers = [
            "Student ID",
            "Name",
            "Enrollment No",
            "Year",
            "Semester",
            "Section",
            "Present Days",
            "Total Days",
            "Attendance %"
        ]

        summary_sheet.merge_cells(
            "A1:I1"
        )

        summary_sheet["A1"] = (
            "AUTOMATED ATTENDANCE REPORT"
        )

        summary_sheet["A1"].font = Font(
            bold=True,
            size=16,
            color="FFFFFF"
        )

        summary_sheet["A1"].fill = (
            PatternFill(
                "solid",
                fgColor="1E3A8A"
            )
        )

        summary_sheet["A1"].alignment = (
            Alignment(
                horizontal="center",
                vertical="center"
            )
        )

        summary_sheet.row_dimensions[
            1
        ].height = 30

        for column, header in enumerate(
            summary_headers,
            start=1
        ):

            cell = summary_sheet.cell(
                row=3,
                column=column
            )

            cell.value = header
            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )

            cell.fill = (
                PatternFill(
                    "solid",
                    fgColor="2563EB"
                )
            )

            cell.alignment = (
                Alignment(
                    horizontal="center"
                )
            )

        row_number = 4

        for student_id, student in (
            students.items()
        ):

            present = present_count.get(
                student_id,
                0
            )

            percentage = (
                present / total_days
            ) * 100

            values = [
                student_id,
                student["name"],
                student["enrollment"],
                student["year"],
                student["semester"],
                student["section"],
                present,
                total_days,
                f"{percentage:.2f}%"
            ]

            for column, value in enumerate(
                values,
                start=1
            ):

                cell = summary_sheet.cell(
                    row=row_number,
                    column=column
                )

                cell.value = value

                cell.alignment = (
                    Alignment(
                        horizontal="center"
                    )
                )

            row_number += 1

        # ==================================================
        # DAILY RECORDS
        # ==================================================

        daily_sheet = (
            workbook.create_sheet(
                "Daily Records"
            )
        )

        daily_headers = (
            "ID",
            "Name",
            "Enrollment No",
            "Year",
            "Semester",
            "Section",
            "Date",
            "Time",
            "Status"
        )

        for column, header in enumerate(
            daily_headers,
            start=1
        ):

            cell = daily_sheet.cell(
                row=1,
                column=column
            )

            cell.value = header

            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )

            cell.fill = (
                PatternFill(
                    "solid",
                    fgColor="2563EB"
                )
            )

            cell.alignment = (
                Alignment(
                    horizontal="center"
                )
            )

        row_number = 2

        for record in records:

            values = [
                record["ID"],
                record["Name"],
                record["Enrollment No"],
                record["Year"],
                record["Semester"],
                record["Section"],
                record["Date"],
                record["Time"],
                record["Status"]
            ]

            for column, value in enumerate(
                values,
                start=1
            ):

                cell = daily_sheet.cell(
                    row=row_number,
                    column=column
                )

                cell.value = value

                cell.alignment = (
                    Alignment(
                        horizontal="center"
                    )
                )

            row_number += 1

        # ==================================================
        # STYLES
        # ==================================================

        border = Border(
            left=Side(
                style="thin",
                color="CBD5E1"
            ),
            right=Side(
                style="thin",
                color="CBD5E1"
            ),
            top=Side(
                style="thin",
                color="CBD5E1"
            ),
            bottom=Side(
                style="thin",
                color="CBD5E1"
            )
        )

        for sheet in (
            summary_sheet,
            daily_sheet
        ):

            for row in sheet.iter_rows():

                for cell in row:
                    cell.border = border

        widths = [
            12,
            24,
            22,
            12,
            14,
            12,
            16,
            14,
            18
        ]

        for index, width in enumerate(
            widths,
            start=1
        ):

            summary_sheet.column_dimensions[
                get_column_letter(index)
            ].width = width

            daily_sheet.column_dimensions[
                get_column_letter(index)
            ].width = width

        summary_sheet.freeze_panes = "A4"
        daily_sheet.freeze_panes = "A2"

        workbook.save(
            REPORT_FILE
        )

        status_label.config(
            text="Excel report exported successfully",
            fg=GREEN
        )

        messagebox.showinfo(
            "Export Successful",
            "Attendance report exported successfully!\n\n"
            f"File:\n{REPORT_FILE}"
        )

    except Exception as error:

        messagebox.showerror(
            "Export Error",
            f"Could not create Excel report.\n\n{error}"
        )


# ==================================================
# BUTTON HELPERS
# ==================================================

def button_enter(button):
    button.config(
        bg=BLUE_HOVER
    )


def button_leave(button):
    button.config(
        bg=BLUE
    )


def create_button(
    parent,
    text,
    command
):

    button = tk.Button(
        parent,
        text=text,
        command=command,
        font=("Segoe UI", 9, "bold"),
        bg=BLUE,
        fg=WHITE,
        activebackground=BLUE_HOVER,
        activeforeground=WHITE,
        relief="flat",
        bd=0,
        cursor="hand2",
        width=19,
        height=2
    )

    button.pack(
        side="left",
        padx=5
    )

    button.bind(
        "<Enter>",
        lambda event:
        button_enter(button)
    )

    button.bind(
        "<Leave>",
        lambda event:
        button_leave(button)
    )

    return button


# ==================================================
# MAIN WINDOW
# ==================================================

root = tk.Tk()

root.title(
    "Attendance Records"
)

root.geometry(
    "1250x760"
)

root.resizable(
    False,
    False
)

root.configure(
    bg=BG
)


# ==================================================
# HEADER
# ==================================================

header = tk.Frame(
    root,
    bg=BG
)

header.pack(
    fill="x",
    padx=40,
    pady=(25, 5)
)

tk.Label(
    header,
    text="Attendance Records",
    font=("Segoe UI", 26, "bold"),
    fg=WHITE,
    bg=BG
).pack(
    anchor="w"
)

tk.Label(
    header,
    text=(
        "View, analyse and export student attendance "
        "by academic group"
    ),
    font=("Segoe UI", 10),
    fg=MUTED,
    bg=BG
).pack(
    anchor="w",
    pady=(3, 0)
)


# ==================================================
# SUMMARY CARD
# ==================================================

summary_card = tk.Frame(
    root,
    bg=CARD,
    height=65
)

summary_card.pack(
    fill="x",
    padx=40,
    pady=18
)

summary_card.pack_propagate(
    False
)

tk.Label(
    summary_card,
    text="📊",
    font=("Segoe UI", 18),
    bg=CARD
).pack(
    side="left",
    padx=(18, 8)
)

summary_label = tk.Label(
    summary_card,
    text="",
    font=("Segoe UI", 10, "bold"),
    fg=LIGHT_TEXT,
    bg=CARD
)

summary_label.pack(
    side="left"
)


# ==================================================
# TABLE
# ==================================================

table_frame = tk.Frame(
    root,
    bg=CARD
)

table_frame.pack(
    fill="both",
    expand=True,
    padx=40
)


style = ttk.Style()

try:
    style.theme_use("clam")
except Exception:
    pass

style.configure(
    "Treeview",
    background=CARD,
    foreground=LIGHT_TEXT,
    fieldbackground=CARD,
    rowheight=32,
    font=("Segoe UI", 9),
    borderwidth=0
)

style.configure(
    "Treeview.Heading",
    background="#334155",
    foreground=WHITE,
    font=("Segoe UI", 9, "bold"),
    relief="flat"
)

style.map(
    "Treeview",
    background=[
        ("selected", BLUE)
    ],
    foreground=[
        ("selected", WHITE)
    ]
)


columns = (
    "ID",
    "Name",
    "Enrollment No",
    "Year",
    "Semester",
    "Section",
    "Present",
    "Total Days",
    "Attendance %"
)

table = ttk.Treeview(
    table_frame,
    columns=columns,
    show="headings",
    height=12
)

column_widths = {
    "ID": 65,
    "Name": 175,
    "Enrollment No": 145,
    "Year": 75,
    "Semester": 85,
    "Section": 75,
    "Present": 80,
    "Total Days": 90,
    "Attendance %": 110
}

for column in columns:

    table.heading(
        column,
        text=column
    )

    table.column(
        column,
        width=column_widths[column],
        anchor="center"
    )

table.pack(
    fill="both",
    expand=True,
    padx=10,
    pady=10
)


# ==================================================
# BUTTONS
# ==================================================

button_area = tk.Frame(
    root,
    bg=BG
)

button_area.pack(
    pady=12
)

row1 = tk.Frame(
    button_area,
    bg=BG
)

row1.pack(
    pady=3
)

create_button(
    row1,
    "📊 SUMMARY",
    calculate_attendance
)

create_button(
    row1,
    "📋 DAILY RECORDS",
    show_daily_attendance
)

create_button(
    row1,
    "📈 ANALYTICS",
    open_analytics
)

row2 = tk.Frame(
    button_area,
    bg=BG
)

row2.pack(
    pady=3
)

create_button(
    row2,
    "📥 EXPORT TO EXCEL",
    export_to_excel
)

create_button(
    row2,
    "🔄 REFRESH",
    calculate_attendance
)


# ==================================================
# STATUS
# ==================================================

status_frame = tk.Frame(
    root,
    bg=BG
)

status_frame.pack(
    pady=(0, 12)
)

tk.Label(
    status_frame,
    text="●",
    font=("Segoe UI", 10),
    fg=GREEN,
    bg=BG
).pack(
    side="left",
    padx=(0, 6)
)

status_label = tk.Label(
    status_frame,
    text="System Ready",
    font=("Segoe UI", 9, "bold"),
    fg=GREEN,
    bg=BG
)

status_label.pack(
    side="left"
)


# ==================================================
# START
# ==================================================

initialize_attendance_excel()
calculate_attendance()

root.mainloop()
