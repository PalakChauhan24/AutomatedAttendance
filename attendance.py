import cv2
import os
import csv
import pandas as pd
from datetime import datetime
from openpyxl import Workbook, load_workbook

try:
    from email_notification import send_student_attendance_email
except ImportError:
    send_student_attendance_email = None


# =========================================================
# SETTINGS
# =========================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

DATA_DIR = os.path.join(BASE_DIR, "data")
ATTENDANCE_DIR = os.path.join(BASE_DIR, "attendance")
TRAINER_DIR = os.path.join(BASE_DIR, "trainer")

STUDENT_FILE = os.path.join(
    DATA_DIR,
    "students.csv"
)

CASCADE_FILE = os.path.join(
    BASE_DIR,
    "haarcascade_frontalface_default.xml"
)

MODEL_FILE = os.path.join(
    TRAINER_DIR,
    "trainingData.yml"
)

EXCEL_FILE = os.path.join(
    ATTENDANCE_DIR,
    "attendance_records.xlsx"
)

THRESHOLD = 102
REQUIRED_FRAMES = 8

STUDENT_HEADERS = [
    "ID",
    "Name",
    "Enrollment No",
    "Email",
    "Year",
    "Semester",
    "Section"
]

ATTENDANCE_HEADERS = [
    "ID",
    "Name",
    "Enrollment No",
    "Year",
    "Semester",
    "Section",
    "Date",
    "Time",
    "Status"
]


# =========================================================
# DIRECTORIES
# =========================================================

os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(ATTENDANCE_DIR, exist_ok=True)
os.makedirs(TRAINER_DIR, exist_ok=True)


# =========================================================
# UPGRADE STUDENTS.CSV
# =========================================================

def prepare_students_csv():

    if not os.path.exists(STUDENT_FILE):

        with open(
            STUDENT_FILE,
            "w",
            newline="",
            encoding="utf-8"
        ) as file:

            csv.writer(file).writerow(
                STUDENT_HEADERS
            )

        return

    try:

        with open(
            STUDENT_FILE,
            "r",
            newline="",
            encoding="utf-8-sig"
        ) as file:

            reader = csv.DictReader(file)
            headers = reader.fieldnames or []
            rows = list(reader)

        if all(
            header in headers
            for header in STUDENT_HEADERS
        ):
            return

        upgraded_rows = []

        for row in rows:

            upgraded_rows.append([
                row.get("ID", ""),
                row.get("Name", ""),
                row.get("Enrollment No", ""),
                row.get("Email", ""),
                row.get("Year", "Not Specified"),
                row.get("Semester", "Not Specified"),
                row.get("Section", "Not Specified")
            ])

        with open(
            STUDENT_FILE,
            "w",
            newline="",
            encoding="utf-8"
        ) as file:

            writer = csv.writer(file)
            writer.writerow(STUDENT_HEADERS)
            writer.writerows(upgraded_rows)

    except Exception as error:

        print(
            "⚠️ Could not upgrade students.csv:",
            error
        )


prepare_students_csv()


# =========================================================
# LOAD STUDENTS
# =========================================================

def load_students():

    students = {}

    if not os.path.exists(STUDENT_FILE):

        print("❌ students.csv not found.")
        return students

    try:

        df = pd.read_csv(
            STUDENT_FILE,
            encoding="utf-8-sig"
        )

    except Exception as error:

        print(
            "❌ Error reading students.csv:",
            error
        )

        return students

    for _, row in df.iterrows():

        try:
            student_id = int(
                str(row["ID"]).strip()
            )
        except Exception:
            continue

        students[student_id] = {
            "name": str(
                row.get("Name", "")
            ).strip(),

            "enrollment": str(
                row.get("Enrollment No", "")
            ).strip(),

            "email": str(
                row.get("Email", "")
            ).strip(),

            "year": str(
                row.get("Year", "Not Specified")
            ).strip(),

            "semester": str(
                row.get("Semester", "Not Specified")
            ).strip(),

            "section": str(
                row.get("Section", "Not Specified")
            ).strip()
        }

    return students


# =========================================================
# CREATE / UPGRADE EXCEL WORKBOOK
# =========================================================

def initialize_excel():

    if not os.path.exists(EXCEL_FILE):

        wb = Workbook()

        ws_attendance = wb.active
        ws_attendance.title = "Attendance Records"

        ws_attendance.append(
            ATTENDANCE_HEADERS
        )

        ws_students = wb.create_sheet(
            "Student Records"
        )

        ws_students.append(
            STUDENT_HEADERS
        )

        wb.save(EXCEL_FILE)

        print(
            "✅ Excel attendance workbook created."
        )

    else:

        upgrade_excel_workbook()

    update_student_records()


# =========================================================
# UPGRADE EXISTING EXCEL
# =========================================================

def upgrade_excel_workbook():

    try:

        wb = load_workbook(EXCEL_FILE)

        # ---------------- ATTENDANCE SHEET ----------------

        if "Attendance Records" not in wb.sheetnames:

            ws = wb.create_sheet(
                "Attendance Records"
            )

            ws.append(
                ATTENDANCE_HEADERS
            )

        else:

            ws = wb["Attendance Records"]

            current_headers = [
                ws.cell(
                    row=1,
                    column=column
                ).value
                for column in range(
                    1,
                    ws.max_column + 1
                )
            ]

            if current_headers != ATTENDANCE_HEADERS:

                old_rows = list(
                    ws.iter_rows(
                        min_row=2,
                        values_only=True
                    )
                )

                wb.remove(ws)

                ws = wb.create_sheet(
                    "Attendance Records",
                    0
                )

                ws.append(
                    ATTENDANCE_HEADERS
                )

                old_index = {
                    str(header): index
                    for index, header in enumerate(
                        current_headers
                    )
                }

                for old_row in old_rows:

                    def old_value(header):
                        index = old_index.get(header)
                        if index is None:
                            return ""
                        if index >= len(old_row):
                            return ""
                        return old_row[index] or ""

                    ws.append([
                        old_value("ID"),
                        old_value("Name"),
                        old_value("Enrollment No"),
                        old_value("Year") or "Not Specified",
                        old_value("Semester") or "Not Specified",
                        old_value("Section") or "Not Specified",
                        old_value("Date"),
                        old_value("Time"),
                        old_value("Status")
                    ])

        # ---------------- STUDENT SHEET ----------------

        if "Student Records" not in wb.sheetnames:

            ws_students = wb.create_sheet(
                "Student Records"
            )

            ws_students.append(
                STUDENT_HEADERS
            )

        else:

            ws_students = wb[
                "Student Records"
            ]

            current_headers = [
                ws_students.cell(
                    row=1,
                    column=column
                ).value
                for column in range(
                    1,
                    ws_students.max_column + 1
                )
            ]

            if current_headers != STUDENT_HEADERS:

                old_rows = list(
                    ws_students.iter_rows(
                        min_row=2,
                        values_only=True
                    )
                )

                wb.remove(ws_students)

                ws_students = wb.create_sheet(
                    "Student Records"
                )

                ws_students.append(
                    STUDENT_HEADERS
                )

                old_index = {
                    str(header): index
                    for index, header in enumerate(
                        current_headers
                    )
                }

                for old_row in old_rows:

                    def old_value(header):
                        index = old_index.get(header)
                        if index is None:
                            return ""
                        if index >= len(old_row):
                            return ""
                        return old_row[index] or ""

                    ws_students.append([
                        old_value("ID"),
                        old_value("Name"),
                        old_value("Enrollment No"),
                        old_value("Email"),
                        old_value("Year") or "Not Specified",
                        old_value("Semester") or "Not Specified",
                        old_value("Section") or "Not Specified"
                    ])

        wb.save(EXCEL_FILE)

    except Exception as error:

        print(
            "⚠️ Could not upgrade Excel workbook:",
            error
        )


# =========================================================
# UPDATE STUDENT RECORDS IN EXCEL
# =========================================================

def update_student_records():

    students = load_students()

    if not students:
        return

    try:

        wb = load_workbook(
            EXCEL_FILE
        )

        if "Student Records" not in wb.sheetnames:

            ws = wb.create_sheet(
                "Student Records"
            )

            ws.append(
                STUDENT_HEADERS
            )

        ws = wb["Student Records"]

        if ws.max_row > 1:

            ws.delete_rows(
                2,
                ws.max_row - 1
            )

        for student_id in sorted(
            students.keys()
        ):

            student = students[
                student_id
            ]

            ws.append([
                student_id,
                student["name"],
                student["enrollment"],
                student["email"],
                student["year"],
                student["semester"],
                student["section"]
            ])

        widths = {
            "A": 10,
            "B": 25,
            "C": 25,
            "D": 35,
            "E": 12,
            "F": 14,
            "G": 12
        }

        for column, width in widths.items():

            ws.column_dimensions[
                column
            ].width = width

        wb.save(
            EXCEL_FILE
        )

        print(
            "✅ Student records synchronized with Excel."
        )

    except Exception as error:

        print(
            "❌ Could not update student records:",
            error
        )


# =========================================================
# CHECK EXCEL DUPLICATE
# =========================================================

def already_marked_in_excel(
    student_id,
    date
):

    if not os.path.exists(EXCEL_FILE):
        return False

    try:

        wb = load_workbook(
            EXCEL_FILE,
            read_only=True
        )

        if "Attendance Records" not in wb.sheetnames:

            wb.close()
            return False

        ws = wb[
            "Attendance Records"
        ]

        headers = {
            str(cell.value): index
            for index, cell in enumerate(
                ws[1],
                start=1
            )
        }

        id_column = headers.get("ID")
        date_column = headers.get("Date")

        if not id_column or not date_column:

            wb.close()
            return False

        for row in ws.iter_rows(
            min_row=2,
            values_only=True
        ):

            existing_id = row[
                id_column - 1
            ]

            existing_date = row[
                date_column - 1
            ]

            try:
                existing_id = int(
                    existing_id
                )
            except Exception:
                continue

            if (
                existing_id == student_id
                and str(existing_date) == date
            ):

                wb.close()
                return True

        wb.close()

    except Exception as error:

        print(
            "⚠️ Excel duplicate check failed:",
            error
        )

    return False


# =========================================================
# SAVE ATTENDANCE TO EXCEL
# =========================================================

def save_attendance_to_excel(
    student_id,
    student,
    date,
    time,
    status
):

    try:

        initialize_excel()

        if already_marked_in_excel(
            student_id,
            date
        ):

            print(
                f"ℹ️ Attendance already exists "
                f"for ID {student_id} on {date}."
            )

            return False

        wb = load_workbook(
            EXCEL_FILE
        )

        ws = wb[
            "Attendance Records"
        ]

        ws.append([
            student_id,
            student["name"],
            student["enrollment"],
            student["year"],
            student["semester"],
            student["section"],
            date,
            time,
            status
        ])

        widths = {
            "A": 10,
            "B": 25,
            "C": 25,
            "D": 12,
            "E": 14,
            "F": 12,
            "G": 15,
            "H": 15,
            "I": 15
        }

        for column, width in widths.items():

            ws.column_dimensions[
                column
            ].width = width

        wb.save(
            EXCEL_FILE
        )

        print(
            "✅ Attendance saved to Excel."
        )

        return True

    except Exception as error:

        print(
            "❌ Could not save attendance to Excel:",
            error
        )

        return False


# =========================================================
# DAILY CSV
# =========================================================

def get_today_csv():

    today = datetime.now().strftime(
        "%Y-%m-%d"
    )

    return os.path.join(
        ATTENDANCE_DIR,
        f"attendance_{today}.csv"
    )


def initialize_daily_csv():

    csv_file = get_today_csv()

    if not os.path.exists(csv_file):

        with open(
            csv_file,
            "w",
            newline="",
            encoding="utf-8-sig"
        ) as file:

            csv.writer(file).writerow([
                "ID",
                "Name",
                "Enrollment No",
                "Year",
                "Semester",
                "Section",
                "Date",
                "Time",
                "Status"
            ])

    return csv_file


# =========================================================
# LOAD TODAY'S MARKED STUDENTS
# =========================================================

def load_marked_students():

    marked_students = set()

    csv_file = get_today_csv()

    if os.path.exists(csv_file):

        try:

            with open(
                csv_file,
                "r",
                newline="",
                encoding="utf-8-sig"
            ) as file:

                reader = csv.DictReader(
                    file
                )

                for row in reader:

                    try:
                        marked_students.add(
                            int(
                                str(
                                    row["ID"]
                                ).strip()
                            )
                        )
                    except Exception:
                        pass

        except Exception as error:

            print(
                "⚠️ Could not read today's CSV:",
                error
            )

    # Also check Excel so duplicate prevention
    # remains active even if today's CSV is missing.
    if os.path.exists(EXCEL_FILE):

        try:

            wb = load_workbook(
                EXCEL_FILE,
                read_only=True
            )

            if "Attendance Records" in wb.sheetnames:

                ws = wb[
                    "Attendance Records"
                ]

                headers = {
                    str(cell.value): index
                    for index, cell in enumerate(
                        ws[1],
                        start=1
                    )
                }

                id_column = headers.get("ID")
                date_column = headers.get("Date")

                if id_column and date_column:

                    today = datetime.now().strftime(
                        "%Y-%m-%d"
                    )

                    for row in ws.iter_rows(
                        min_row=2,
                        values_only=True
                    ):

                        try:

                            row_id = int(
                                row[
                                    id_column - 1
                                ]
                            )

                            row_date = str(
                                row[
                                    date_column - 1
                                ]
                            )

                            if row_date == today:
                                marked_students.add(
                                    row_id
                                )

                        except Exception:
                            continue

            wb.close()

        except Exception as error:

            print(
                "⚠️ Could not check Excel attendance:",
                error
            )

    return marked_students


# =========================================================
# SAVE ATTENDANCE TO CSV
# =========================================================

def save_attendance_to_csv(
    student_id,
    student,
    date,
    time,
    status
):

    csv_file = initialize_daily_csv()

    with open(
        csv_file,
        "a",
        newline="",
        encoding="utf-8-sig"
    ) as file:

        csv.writer(file).writerow([
            student_id,
            student["name"],
            student["enrollment"],
            student["year"],
            student["semester"],
            student["section"],
            date,
            time,
            status
        ])


# =========================================================
# MARK ATTENDANCE
# =========================================================

def mark_attendance(
    student_id,
    students,
    marked_students
):

    if student_id in marked_students:

        print(
            f"ℹ️ ID {student_id} already marked present today."
        )

        return

    if student_id not in students:

        print(
            f"⚠️ Student ID {student_id} not found."
        )

        return

    student = students[
        student_id
    ]

    now = datetime.now()

    date = now.strftime(
        "%Y-%m-%d"
    )

    time = now.strftime(
        "%H:%M:%S"
    )

    # Save to Excel first.
    excel_saved = save_attendance_to_excel(
        student_id,
        student,
        date,
        time,
        "PRESENT"
    )

    if not excel_saved:

        print(
            "⚠️ Attendance was not added to Excel."
        )

        return

    # Keep daily CSV for compatibility.
    save_attendance_to_csv(
        student_id,
        student,
        date,
        time,
        "PRESENT"
    )

    marked_students.add(
        student_id
    )

    print()
    print("========================================")
    print("        ATTENDANCE MARKED")
    print("========================================")
    print(f"ID          : {student_id}")
    print(f"Name        : {student['name']}")
    print(f"Enrollment  : {student['enrollment']}")
    print(f"Year        : {student['year']}")
    print(f"Semester    : {student['semester']}")
    print(f"Section     : {student['section']}")
    print(f"Date        : {date}")
    print(f"Time        : {time}")
    print("Status      : PRESENT")
    print("Excel Saved : YES")
    print("========================================")
    print()

    # Email notification
    if send_student_attendance_email:

        try:

            send_student_attendance_email(
                student_id
            )

        except Exception as error:

            print(
                "⚠️ Email notification failed:",
                error
            )


# =========================================================
# MAIN FACE RECOGNITION
# =========================================================

def start_attendance():

    print()
    print("========================================")
    print("      AUTOMATED ATTENDANCE SYSTEM")
    print("========================================")
    print()

    students = load_students()

    if not students:

        print(
            "❌ No students registered."
        )

        return

    print(
        f"✅ Registered students: "
        f"{len(students)}"
    )

    initialize_excel()
    initialize_daily_csv()

    marked_students = load_marked_students()

    print(
        f"✅ Already marked today: "
        f"{len(marked_students)}"
    )

    # -----------------------------------------------------
    # Haar Cascade
    # -----------------------------------------------------

    face_cascade = cv2.CascadeClassifier(
        CASCADE_FILE
    )

    if face_cascade.empty():

        print(
            "❌ Could not load Haar Cascade."
        )

        return

    # -----------------------------------------------------
    # LBPH Model
    # -----------------------------------------------------

    if not os.path.exists(MODEL_FILE):

        print(
            "❌ Training model not found."
        )

        print(
            "👉 Please run train.py first."
        )

        return

    recognizer = (
        cv2.face.LBPHFaceRecognizer_create()
    )

    try:

        recognizer.read(
            MODEL_FILE
        )

    except Exception as error:

        print(
            "❌ Could not load face model:",
            error
        )

        return

    # -----------------------------------------------------
    # Camera
    # -----------------------------------------------------

    camera = cv2.VideoCapture(0)

    if not camera.isOpened():

        print(
            "❌ Could not open camera."
        )

        return

    print()
    print("📷 Camera started.")
    print("👉 Look at the camera.")
    print("👉 Press Q to quit.")
    print()

    last_detected_id = None
    consecutive_count = 0

    # =====================================================
    # CAMERA LOOP
    # =====================================================

    while True:

        ret, frame = camera.read()

        if not ret:

            print(
                "⚠️ Could not read camera frame."
            )

            break

        # -------------------------------------------------
        # Grayscale
        # -------------------------------------------------

        gray = cv2.cvtColor(
            frame,
            cv2.COLOR_BGR2GRAY
        )

        # -------------------------------------------------
        # Preprocessing for Haar detection
        # -------------------------------------------------

        equalized = cv2.equalizeHist(
            gray
        )

        processed = cv2.medianBlur(
            equalized,
            5
        )

        # -------------------------------------------------
        # Detect faces
        # -------------------------------------------------

        faces = face_cascade.detectMultiScale(
            processed,
            scaleFactor=1.2,
            minNeighbors=5,
            minSize=(100, 100)
        )

        recognized_this_frame = False

        for (x, y, w, h) in faces:

            # -------------------------------------------------
            # Original grayscale ROI for LBPH
            # -------------------------------------------------

            face_roi = gray[
                y:y + h,
                x:x + w
            ]

            if face_roi.size == 0:
                continue

            face_roi = cv2.resize(
                face_roi,
                (100, 100)
            )

            try:

                student_id, distance = (
                    recognizer.predict(
                        face_roi
                    )
                )

                student_id = int(
                    student_id
                )

                distance = float(
                    distance
                )

            except Exception as error:

                print(
                    "⚠️ Recognition error:",
                    error
                )

                continue

            # -------------------------------------------------
            # Valid recognition
            # -------------------------------------------------

            if (
                distance <= THRESHOLD
                and student_id in students
            ):

                recognized_this_frame = True

                student = students[
                    student_id
                ]

                name = student[
                    "name"
                ]

                label = (
                    f"{name} "
                    f"({distance:.2f})"
                )

                box_color = (
                    0,
                    255,
                    0
                )

                # Consecutive confirmation
                if (
                    student_id
                    == last_detected_id
                ):

                    consecutive_count += 1

                else:

                    last_detected_id = (
                        student_id
                    )

                    consecutive_count = 1

                cv2.putText(
                    frame,
                    (
                        f"Confirming: "
                        f"{consecutive_count}/"
                        f"{REQUIRED_FRAMES}"
                    ),
                    (x, y - 35),
                    cv2.FONT_HERSHEY_SIMPLEX,
                    0.7,
                    box_color,
                    2
                )

                if (
                    consecutive_count
                    >= REQUIRED_FRAMES
                ):

                    mark_attendance(
                        student_id,
                        students,
                        marked_students
                    )

                    last_detected_id = None
                    consecutive_count = 0

            else:

                label = (
                    "Unknown Student"
                )

                box_color = (
                    0,
                    0,
                    255
                )

        # Reset confirmation if no valid
        # recognition was present in this frame.
        if not recognized_this_frame:

            last_detected_id = None
            consecutive_count = 0

        # -------------------------------------------------
        # Draw face information
        # -------------------------------------------------

        for (x, y, w, h) in faces:

            try:

                current_face = gray[
                    y:y + h,
                    x:x + w
                ]

                if current_face.size == 0:
                    continue

                current_face = cv2.resize(
                    current_face,
                    (100, 100)
                )

                current_id, current_distance = (
                    recognizer.predict(
                        current_face
                    )
                )

                current_id = int(
                    current_id
                )

                current_distance = float(
                    current_distance
                )

                if (
                    current_distance <= THRESHOLD
                    and current_id in students
                ):

                    display_label = (
                        f"{students[current_id]['name']} "
                        f"({current_distance:.2f})"
                    )

                    display_color = (
                        0,
                        255,
                        0
                    )

                else:

                    display_label = (
                        "Unknown Student"
                    )

                    display_color = (
                        0,
                        0,
                        255
                    )

                cv2.rectangle(
                    frame,
                    (x, y),
                    (x + w, y + h),
                    display_color,
                    2
                )

                cv2.putText(
                    frame,
                    display_label,
                    (x, y + h + 25),
                    cv2.FONT_HERSHEY_SIMPLEX,
                    0.7,
                    display_color,
                    2
                )

            except Exception:
                continue

        # -------------------------------------------------
        # System information
        # -------------------------------------------------

        cv2.putText(
            frame,
            "Automated Attendance System",
            (20, 35),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.8,
            (255, 255, 255),
            2
        )

        cv2.putText(
            frame,
            "Press Q to Exit",
            (20, 65),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.6,
            (255, 255, 255),
            2
        )

        cv2.imshow(
            "Automated Attendance",
            frame
        )

        key = cv2.waitKey(1) & 0xFF

        if key == ord("q"):
            break

    # =====================================================
    # CLEANUP
    # =====================================================

    camera.release()

    cv2.destroyAllWindows()

    print()
    print("📷 Camera stopped.")
    print("✅ Attendance session ended.")


# =========================================================
# RUN
# =========================================================

if __name__ == "__main__":
    start_attendance()
